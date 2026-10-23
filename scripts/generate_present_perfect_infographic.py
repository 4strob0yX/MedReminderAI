import os
import shutil
from PIL import Image, ImageDraw, ImageFont, ImageEnhance
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# ==============================================================================
# COLOR PALETTE & TYPOGRAPHY - MODERN COMMERCIAL HEALTHCARE
# ==============================================================================
FONT_PRIMARY = "Helvetica Neue"

HEX_PRIMARY_DARK  = "0F172A"  # Slate 900 (Deep Midnight Navy)
HEX_PRIMARY_ACCENT= "0284C7"  # Sky 600 (Medical Cyan-Blue)
HEX_ACCENT_SLATE  = "334155"  # Slate 700
HEX_BODY_TEXT     = "1E293B"  # Slate 800
HEX_MUTED_TEXT    = "64748B"  # Slate 500
HEX_BORDER_LIGHT  = "E2E8F0"  # Slate 200 (Clean Hairline)
HEX_SURFACE_SOFT  = "F8FAFC"  # Slate 50
HEX_SURFACE_ALT   = "F1F5F9"  # Slate 100
HEX_SUCCESS_GREEN = "059669"  # Emerald 600
HEX_VERB_BLUE     = "0369A1"  # Sky 700 (High-Contrast Verb Highlight)

RGB_DARK   = RGBColor(15, 23, 42)     # #0F172A
RGB_ACCENT = RGBColor(2, 132, 199)    # #0284C7
RGB_BODY   = RGBColor(30, 41, 59)     # #1E293B
RGB_MUTED  = RGBColor(100, 116, 139)  # #64748B
RGB_LIGHT  = RGBColor(226, 232, 240)  # #E2E8F0
RGB_VERB   = RGBColor(3, 105, 161)    # #0369A1

# ==============================================================================
# 15 CONSUMER SALES PITCH VALUE PROPOSITIONS (100% NON-TECHNICAL, ULTRA-SHORT)
# Step 01: Explains the name clearly (Med = Medicine, Reminder = Reminders, AI = Smart Care)
# Pure marketing language for everyday customers, families, and buyers.
# 15 Unique Action Verbs in Present Perfect (have / has + V3).
# ==============================================================================
VERBS_PRESENT_PERFECT_DATA = [
    # --------------------------------------------------------------------------
    # PILLAR 01: NAME ORIGIN & ACCESSIBILITY (STEPS 01–05)
    # --------------------------------------------------------------------------
    {
        "num": "01",
        "phase": 1,
        "title": "Why We Chose MedReminder AI",
        "base_verb": "choose",
        "participle": "chosen",
        "verb_type": "Irregular",
        "auxiliary": "have",
        "perfect_phrase": "have chosen",
        "pre_text": "We ",
        "post_text": " MedReminder AI for Medicine, Reminders, and Smart Care.",
        "sales_pitch": "A clear name that combines your medicine, punctual reminders, and smart care.",
        "tech_tag": "Brand Meaning  ·  Medicine + Reminder + AI",
        "spanish_translation": "Elegimos MedReminder AI por Medicina (Med), Recordatorios (Reminder) y Cuidado Inteligente (AI).",
        "benefit_es": "Nombre claro y confiable: medicinas a tiempo con ayuda inteligente."
    },
    {
        "num": "02",
        "phase": 1,
        "title": "Big Buttons for Seniors",
        "base_verb": "design",
        "participle": "designed",
        "verb_type": "Regular",
        "auxiliary": "have",
        "perfect_phrase": "have designed",
        "pre_text": "We ",
        "post_text": " big, clear buttons for seniors.",
        "sales_pitch": "Large screen buttons with zero accidental clicks for elderly family members.",
        "tech_tag": "Easy to Use  ·  Big Screen Buttons",
        "spanish_translation": "Hemos diseñado botones grandes y claros para adultos mayores.",
        "benefit_es": "Letras y botones gigantes: cero confusiones para tus papás o abuelos."
    },
    {
        "num": "03",
        "phase": 1,
        "title": "Zero Double Doses",
        "base_verb": "eliminate",
        "participle": "eliminated",
        "verb_type": "Regular",
        "auxiliary": "has",
        "perfect_phrase": "has eliminated",
        "pre_text": "Our app ",
        "post_text": " accidental double doses.",
        "sales_pitch": "Guarantees you never take the same medicine twice by mistake.",
        "tech_tag": "Total Safety  ·  Never Double Take",
        "spanish_translation": "Nuestra app ha eliminado las tomas dobles por accidente.",
        "benefit_es": "Seguridad total: te protege para que nunca repitas una pastilla por error."
    },
    {
        "num": "04",
        "phase": 1,
        "title": "Alarms Without Internet",
        "base_verb": "build",
        "participle": "built",
        "verb_type": "Irregular",
        "auxiliary": "have",
        "perfect_phrase": "have built",
        "pre_text": "We ",
        "post_text": " alarms that ring without internet.",
        "sales_pitch": "Punctual alerts anywhere: in the basement, traveling, or on airplanes.",
        "tech_tag": "Always Rings  ·  100% Offline",
        "spanish_translation": "Hemos construido alarmas que suenan incluso sin internet.",
        "benefit_es": "Tus alarmas suenan siempre a tiempo, con o sin señal o datos móviles."
    },
    {
        "num": "05",
        "phase": 1,
        "title": "Digital Pillbox",
        "base_verb": "simplify",
        "participle": "simplified",
        "verb_type": "Regular",
        "auxiliary": "has",
        "perfect_phrase": "has simplified",
        "pre_text": "Our digital box ",
        "post_text": " your family's pills.",
        "sales_pitch": "Organizes your home pharmacy and daily treatments in a single tap.",
        "tech_tag": "Home Pillbox  ·  One-Tap Setup",
        "spanish_translation": "Nuestro pastillero digital ha simplificado las medicinas de tu familia.",
        "benefit_es": "Todo tu botiquín ordenado en segundos sin enredos de horarios."
    },

    # --------------------------------------------------------------------------
    # PILLAR 02: FAMILY PRIVACY & TIMELY ALERTS (STEPS 06–10)
    # --------------------------------------------------------------------------
    {
        "num": "06",
        "phase": 2,
        "title": "Private Health Data",
        "base_verb": "protect",
        "participle": "protected",
        "verb_type": "Regular",
        "auxiliary": "have",
        "perfect_phrase": "have protected",
        "pre_text": "We ",
        "post_text": " your family's private health info.",
        "sales_pitch": "Bank-grade privacy: your sensitive medical life stays strictly confidential.",
        "tech_tag": "Total Privacy  ·  Safe Health Info",
        "spanish_translation": "Hemos protegido la información médica privada de tu familia.",
        "benefit_es": "Privacidad total: las recetas y salud de tu familia están 100% seguras."
    },
    {
        "num": "07",
        "phase": 2,
        "title": "100% Reliability",
        "base_verb": "achieve",
        "participle": "achieved",
        "verb_type": "Regular",
        "auxiliary": "has",
        "perfect_phrase": "has achieved",
        "pre_text": "The app ",
        "post_text": " 100% reliability for our users.",
        "sales_pitch": "Never freezes, never crashes: ready every single second you need your pill.",
        "tech_tag": "Zero Crashes  ·  Always Working",
        "spanish_translation": "La app ha alcanzado un 100% de confiabilidad para nuestros usuarios.",
        "benefit_es": "La app nunca se traba ni se cae; siempre lista cuando la necesitas."
    },
    {
        "num": "08",
        "phase": 2,
        "title": "Refill Reminders",
        "base_verb": "prevent",
        "participle": "prevented",
        "verb_type": "Regular",
        "auxiliary": "has",
        "perfect_phrase": "has prevented",
        "pre_text": "The app ",
        "post_text": " running out of essential medicine.",
        "sales_pitch": "Warns you days in advance so you can buy more pills without rushing.",
        "tech_tag": "Refill Alerts  ·  Never Run Out",
        "spanish_translation": "La app ha prevenido que te quedes sin medicamentos esenciales.",
        "benefit_es": "Te avisa con días de anticipación antes de que se acabe tu medicina."
    },
    {
        "num": "09",
        "phase": 2,
        "title": "Gentle Reminders",
        "base_verb": "introduce",
        "participle": "introduced",
        "verb_type": "Regular",
        "auxiliary": "has",
        "perfect_phrase": "has introduced",
        "pre_text": "Our app ",
        "post_text": " smart reminders to avoid late doses.",
        "sales_pitch": "Gently keeps you on track without letting you procrastinate or forget.",
        "tech_tag": "On-Time Routine  ·  Healthy Habits",
        "spanish_translation": "Nuestra app ha introducido recordatorios inteligentes para evitar retrasos.",
        "benefit_es": "Te cuida con amabilidad para que no dejes pasar la hora de tu tratamiento."
    },
    {
        "num": "10",
        "phase": 2,
        "title": "No Pill Confusion",
        "base_verb": "solve",
        "participle": "solved",
        "verb_type": "Regular",
        "auxiliary": "has",
        "perfect_phrase": "has solved",
        "pre_text": "Our app ",
        "post_text": " pill confusion for busy families.",
        "sales_pitch": "Eliminates morning and nighttime confusion between multiple medications.",
        "tech_tag": "Clear Schedules  ·  Zero Confusion",
        "spanish_translation": "Nuestra app ha resuelto la confusión de pastillas en familias ocupadas.",
        "benefit_es": "Se acabaron las dudas de qué pastilla toca en la mañana o en la noche."
    },

    # --------------------------------------------------------------------------
    # PILLAR 03: FAMILY TRUST & PEACE OF MIND (STEPS 11–15)
    # --------------------------------------------------------------------------
    {
        "num": "11",
        "phase": 3,
        "title": "Instant Notifications",
        "base_verb": "deliver",
        "participle": "delivered",
        "verb_type": "Regular",
        "auxiliary": "has",
        "perfect_phrase": "has delivered",
        "pre_text": "The app ",
        "post_text": " instant alerts right on time.",
        "sales_pitch": "Ultra-responsive notifications that pop up right on the exact second.",
        "tech_tag": "Instant Speed  ·  Alerts on the Second",
        "spanish_translation": "La app ha entregado alertas instantáneas justo a tiempo.",
        "benefit_es": "Rapidez total: suena exactamente al segundo sin ninguna demora."
    },
    {
        "num": "12",
        "phase": 3,
        "title": "Doctor Visit Reports",
        "base_verb": "save",
        "participle": "saved",
        "verb_type": "Regular",
        "auxiliary": "has",
        "perfect_phrase": "has saved",
        "pre_text": "Our app ",
        "post_text": " your complete history for doctor visits.",
        "sales_pitch": "Show your physician clear proof of your daily adherence habits.",
        "tech_tag": "Doctor Reports  ·  Clear History",
        "spanish_translation": "Nuestra app ha guardado tu historial completo para las visitas al doctor.",
        "benefit_es": "Muéstrale a tu médico exactamente cómo y cuándo tomaste cada medicina."
    },
    {
        "num": "13",
        "phase": 3,
        "title": "Peace of Mind",
        "base_verb": "empower",
        "participle": "empowered",
        "verb_type": "Regular",
        "auxiliary": "has",
        "perfect_phrase": "has empowered",
        "pre_text": "MedReminder AI ",
        "post_text": " patients to live with peace of mind.",
        "sales_pitch": "Restores your daily independence so you never stress about medicine clocks.",
        "tech_tag": "Daily Independence  ·  Zero Worry",
        "spanish_translation": "MedReminder AI ha empoderado a los pacientes para vivir tranquilos.",
        "benefit_es": "Te devuelve tu tranquilidad diaria para disfrutar sin estrés de horarios."
    },
    {
        "num": "14",
        "phase": 3,
        "title": "Trusted by Families",
        "base_verb": "trust",
        "participle": "trusted",
        "verb_type": "Regular",
        "auxiliary": "have",
        "perfect_phrase": "have trusted",
        "pre_text": "Thousands of families ",
        "post_text": " our app with their loved ones.",
        "sales_pitch": "The #1 choice for caring sons, daughters, and parents nationwide.",
        "tech_tag": "Family Care  ·  Loved Ones First",
        "spanish_translation": "Miles de familias han confiado en nuestra app para cuidar a sus seres queridos.",
        "benefit_es": "La app favorita de hijos y padres para proteger la salud familiar."
    },
    {
        "num": "15",
        "phase": 3,
        "title": "Saving Lives Daily",
        "base_verb": "prove",
        "participle": "proven",
        "verb_type": "Irregular",
        "auxiliary": "has",
        "perfect_phrase": "has proven",
        "pre_text": "MedReminder AI ",
        "post_text": " that taking pills on time saves lives.",
        "sales_pitch": "The best investment in life: protecting health, avoiding relapses, and living longer.",
        "tech_tag": "Saves Lives  ·  Best Health Investment",
        "spanish_translation": "MedReminder AI ha demostrado que tomar tus medicinas a tiempo salva vidas.",
        "benefit_es": "La mejor inversión: salud diaria, menos recaídas y más vida."
    },
]


# ==============================================================================
# FONT LOADER HELPER (PIL)
# ==============================================================================
def load_font(size, bold=False):
    candidates = [
        ("/System/Library/Fonts/HelveticaNeue.ttc", 1 if bold else 0),
        ("/System/Library/Fonts/Helvetica.ttc", 1 if bold else 0),
        ("/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf", 0),
    ]
    for path, index in candidates:
        if os.path.exists(path):
            try:
                return ImageFont.truetype(path, size, index=index)
            except Exception:
                try:
                    return ImageFont.truetype(path, size)
                except Exception:
                    pass
    return ImageFont.load_default()


# ==============================================================================
# VISUAL ASSETS GENERATION (HIGH-RES 300 DPI GRAPHICS - NON-TECHNICAL)
# ==============================================================================
def generate_visual_assets(screenshots_dir, out_assets_dir):
    os.makedirs(out_assets_dir, exist_ok=True)

    # --------------------------------------------------------------------------
    # 1. 3-STAGE COMMERCIAL PILLARS FLOW BANNER (2100 x 170)
    # --------------------------------------------------------------------------
    flow_w, flow_h = 2100, 170
    flow = Image.new("RGB", (flow_w, flow_h), (255, 255, 255))
    fdraw = ImageDraw.Draw(flow)

    stages = [
        (
            "PILLAR 01  ·  STEPS 01–05",
            "NAME MEANING & EASY USE",
            "Verbs: chosen, designed, eliminated, built, simplified",
            "Focus: MedReminder AI Meaning, Big Senior Buttons & Offline Alarms"
        ),
        (
            "PILLAR 02  ·  STEPS 06–10",
            "PRIVACY & RELIABLE REMINDERS",
            "Verbs: protected, achieved, prevented, introduced, solved",
            "Focus: 100% Private Health Info, Zero Crashes & Refill Alerts"
        ),
        (
            "PILLAR 03  ·  STEPS 11–15",
            "FAMILY TRUST & PEACE OF MIND",
            "Verbs: delivered, saved, empowered, trusted, proven",
            "Focus: Instant Alerts, Doctor Reports & Peace of Mind for Loved Ones"
        ),
    ]

    box_w = 666
    box_h = 152
    for i, (p_lbl, p_main, p_verbs, p_pitch) in enumerate(stages):
        bx = 14 + i * (box_w + 35)
        by = 9
        fdraw.rounded_rectangle([bx, by, bx + box_w, by + box_h], radius=8, fill=(248, 250, 252), outline=(226, 232, 240), width=2)
        # Left blue accent strip
        fdraw.rounded_rectangle([bx, by, bx + 8, by + box_h], radius=3, fill=(2, 132, 199))

        fdraw.text((bx + 26, by + 12), p_lbl, fill=(2, 132, 199), font=load_font(15, bold=True))
        fdraw.text((bx + 26, by + 34), p_main, fill=(15, 23, 42), font=load_font(20, bold=True))
        fdraw.line([(bx + 26, by + 62), (bx + box_w - 26, by + 62)], fill=(226, 232, 240), width=2)
        fdraw.text((bx + 26, by + 74), p_verbs, fill=(51, 65, 85), font=load_font(16, bold=False))
        fdraw.text((bx + 26, by + 114), p_pitch, fill=(15, 23, 42), font=load_font(17, bold=True))

        if i < 2:
            ax = bx + box_w + 9
            ay = by + box_h // 2
            fdraw.line([(ax, ay), (ax + 17, ay)], fill=(148, 163, 184), width=3)
            fdraw.line([(ax + 9, ay - 7), (ax + 17, ay), (ax + 9, ay + 7)], fill=(148, 163, 184), width=3)

    flow_path = os.path.join(out_assets_dir, "pillars_flow_banner.png")
    flow.save(flow_path, dpi=(300, 300))

    # --------------------------------------------------------------------------
    # 2. COMMERCIAL PRODUCT INTERFACE SHOWCASE (2100 x 440)
    # --------------------------------------------------------------------------
    banner_w, banner_h = 2100, 440
    banner = Image.new("RGB", (banner_w, banner_h), (248, 250, 252))
    draw = ImageDraw.Draw(banner)

    draw.rectangle([0, 0, banner_w - 1, banner_h - 1], outline=(226, 232, 240), width=3)
    draw.rectangle([0, 0, banner_w, 7], fill=(2, 132, 199))

    draw.text((40, 16), "COMMERCIAL PRODUCT INTERFACE & SELLING POINTS", fill=(2, 132, 199), font=load_font(16, bold=True))
    draw.text((40, 38), "Big Senior Buttons, Offline Alarms, Private Health Data & Family Peace of Mind", fill=(15, 23, 42), font=load_font(24, bold=True))

    screens = [
        ("01_login.png", "01 / Private Health Portal", "100% Protected Data (Step 06)"),
        ("02_hoy.png", "02 / Today's Treatment", "Alarms Without Internet (Step 04)"),
        ("03_nuevo_medicamento.png", "03 / Simple Setup", "Big Buttons for Seniors (Step 02)"),
        ("04_botiquin.png", "04 / Family Pillbox", "Never Run Out of Pills (Step 08)"),
        ("07_salud_apoyo.png", "05 / Family Peace of Mind", "Trusted by Loved Ones (Step 14)"),
    ]

    card_w = 384
    card_h = 340
    start_x = 40
    gap = 25
    top_y = 78

    for idx, (fname, label, sublabel) in enumerate(screens):
        x = start_x + idx * (card_w + gap)
        y = top_y

        draw.rounded_rectangle([x, y, x + card_w, y + card_h], radius=8, fill=(255, 255, 255), outline=(226, 232, 240), width=2)
        draw.text((x + 20, y + 12), label, fill=(15, 23, 42), font=load_font(18, bold=True))
        draw.text((x + 20, y + 35), sublabel, fill=(100, 116, 139), font=load_font(15, bold=False))
        draw.line([(x + 20, y + 58), (x + card_w - 20, y + 58)], fill=(241, 245, 249), width=2)

        img_path = os.path.join(screenshots_dir, fname)
        if not os.path.exists(img_path):
            img_path = os.path.join(screenshots_dir, "05_analisis_dfs.png")

        if os.path.exists(img_path):
            im = Image.open(img_path).convert("RGB")
            target_h = 260
            ratio = target_h / im.height
            target_w = int(im.width * ratio)
            im_resized = im.resize((target_w, target_h), Image.Resampling.LANCZOS)
            paste_x = x + (card_w - target_w) // 2
            paste_y = y + 68
            draw.rounded_rectangle(
                [paste_x - 2, paste_y - 2, paste_x + target_w + 2, paste_y + target_h + 2],
                radius=6,
                outline=(203, 213, 225),
                width=2,
            )
            banner.paste(im_resized, (paste_x, paste_y))

    showcase_path = os.path.join(out_assets_dir, "product_showcase_banner.png")
    banner.save(showcase_path, dpi=(300, 300))

    return showcase_path, flow_path


# ==============================================================================
# DOCX XML & FORMATTING HELPERS
# ==============================================================================
def configure_modern_word_settings(doc):
    settings = doc.settings.element
    compat = settings.find(qn("w:compat"))
    if compat is None:
        compat = OxmlElement("w:compat")
        settings.append(compat)
    else:
        for child in list(compat):
            if child.get(qn("w:name")) == "compatibilityMode":
                compat.remove(child)

    compat_setting = parse_xml(
        f'<w:compatSetting {nsdecls("w")} w:name="compatibilityMode" w:uri="http://schemas.microsoft.com/office/word" w:val="15"/>'
    )
    compat.append(compat_setting)

    if settings.find(qn("w:hideSpellingErrors")) is None:
        settings.append(parse_xml(f'<w:hideSpellingErrors {nsdecls("w")}/>'))
    if settings.find(qn("w:hideGrammaticalErrors")) is None:
        settings.append(parse_xml(f'<w:hideGrammaticalErrors {nsdecls("w")}/>'))


def force_table_col_widths(table, col_widths_inches):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    total_dxa = int(sum(col_widths_inches) * 1440)
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is not None:
        tblPr.remove(tblW)
    tblPr.append(parse_xml(f'<w:tblW {nsdecls("w")} w:w="{total_dxa}" w:type="dxa"/>'))

    tblLayout = tblPr.find(qn("w:tblLayout"))
    if tblLayout is not None:
        tblPr.remove(tblLayout)
    tblPr.append(parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'))

    tblGrid = table._tbl.find(qn("w:tblGrid"))
    if tblGrid is not None:
        table._tbl.remove(tblGrid)
    new_grid = OxmlElement("w:tblGrid")
    for w_in in col_widths_inches:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(int(w_in * 1440)))
        new_grid.append(gc)
    table._tbl.insert(1, new_grid)

    for row in table.rows:
        for i, w_in in enumerate(col_widths_inches):
            cell = row.cells[i]
            cell.width = Inches(w_in)


def set_explicit_cell_dxa(cell, width_inches):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is not None:
        tcPr.remove(tcW)
    tcPr.append(parse_xml(f'<w:tcW {nsdecls("w")} w:w="{int(width_inches * 1440)}" w:type="dxa"/>'))


def add_micro_spacer(doc, pt_size=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p.paragraph_format.line_spacing = Pt(pt_size)
    r = p.add_run()
    r.font.size = Pt(1)
    return p


def set_cell_margins(cell, top=50, bottom=50, left=95, right=95):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.find(qn("w:tcMar"))
    if tcMar is not None:
        tcPr.remove(tcMar)
    tcMar = OxmlElement("w:tcMar")
    for m, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is not None:
        tcPr.remove(shd)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def set_card_borders(cell, left_accent_hex=HEX_PRIMARY_ACCENT, border_color=HEX_BORDER_LIGHT):
    tcPr = cell._tc.get_or_add_tcPr()
    old_b = tcPr.find(qn("w:tcBorders"))
    if old_b is not None:
        tcPr.remove(old_b)
    borders = parse_xml(
        f"""
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>
            <w:left w:val="single" w:sz="18" w:space="0" w:color="{left_accent_hex}"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>
        </w:tcBorders>
    """
    )
    tcPr.append(borders)


def set_hairline_box_borders(cell, border_color=HEX_BORDER_LIGHT, bottom_accent=None):
    b_col = bottom_accent if bottom_accent else border_color
    b_sz = "12" if bottom_accent else "4"
    tcPr = cell._tc.get_or_add_tcPr()
    old_b = tcPr.find(qn("w:tcBorders"))
    if old_b is not None:
        tcPr.remove(old_b)
    borders = parse_xml(
        f"""
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>
            <w:bottom w:val="single" w:sz="{b_sz}" w:space="0" w:color="{b_col}"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>
        </w:tcBorders>
    """
    )
    tcPr.append(borders)


def set_minimal_table_borders(table, border_color=HEX_BORDER_LIGHT):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f"""
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="8" w:space="0" w:color="{HEX_PRIMARY_DARK}"/>
            <w:bottom w:val="single" w:sz="8" w:space="0" w:color="{HEX_PRIMARY_DARK}"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>
            <w:insideV w:val="none"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
        </w:tblBorders>
    """
    )
    tblPr.append(borders)


def prevent_row_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))


# ==============================================================================
# CARD BUILDER (STEP CELL - 100% COMMERCIAL & USER-FRIENDLY)
# ==============================================================================
def populate_step_cell(cell, step):
    set_cell_background(cell, "FFFFFF")
    set_card_borders(cell, left_accent_hex=HEX_PRIMARY_ACCENT, border_color=HEX_BORDER_LIGHT)
    set_cell_margins(cell, top=45, bottom=45, left=90, right=85)

    # Line 1: STEP NUMBER + TITLE
    p1 = cell.paragraphs[0]
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(1.5)
    p1.paragraph_format.line_spacing = 1.04

    r_num = p1.add_run(f"STEP {step['num']}   ")
    r_num.bold = True
    r_num.font.name = FONT_PRIMARY
    r_num.font.size = Pt(8.5)
    r_num.font.color.rgb = RGB_ACCENT

    r_title = p1.add_run(step["title"])
    r_title.bold = True
    r_title.font.name = FONT_PRIMARY
    r_title.font.size = Pt(9.5)
    r_title.font.color.rgb = RGB_DARK

    # Line 2: GRAMMAR BADGE & USER BENEFIT TAG
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(2)
    p2.paragraph_format.line_spacing = 1.04

    r_badge = p2.add_run(f"VERB #{step['num']}: {step['perfect_phrase'].upper()} (to {step['base_verb']})")
    r_badge.bold = True
    r_badge.font.name = FONT_PRIMARY
    r_badge.font.size = Pt(7)
    r_badge.font.color.rgb = RGB_VERB

    r_tech = p2.add_run(f"   |   {step['tech_tag']}")
    r_tech.font.name = FONT_PRIMARY
    r_tech.font.size = Pt(7)
    r_tech.font.color.rgb = RGB_MUTED

    # Line 3: ULTRA-SHORT CONSUMER PITCH (ENGLISH)
    p3 = cell.add_paragraph()
    p3.paragraph_format.space_before = Pt(0)
    p3.paragraph_format.space_after = Pt(2.5)
    p3.paragraph_format.line_spacing = 1.08

    r_pre = p3.add_run(step["pre_text"])
    r_pre.font.name = FONT_PRIMARY
    r_pre.font.size = Pt(9.5)
    r_pre.font.color.rgb = RGB_BODY

    # Highlighted Present Perfect phrase
    r_perf = p3.add_run(step["perfect_phrase"])
    r_perf.bold = True
    r_perf.underline = True
    r_perf.font.name = FONT_PRIMARY
    r_perf.font.size = Pt(9.5)
    r_perf.font.color.rgb = RGB_VERB

    r_post = p3.add_run(step["post_text"])
    r_post.font.name = FONT_PRIMARY
    r_post.font.size = Pt(9.5)
    r_post.font.color.rgb = RGB_BODY

    # Line 4: SPANISH TRANSLATION (CLEAR & ACCESSIBLE)
    p4 = cell.add_paragraph()
    p4.paragraph_format.space_before = Pt(0)
    p4.paragraph_format.space_after = Pt(1.5)
    p4.paragraph_format.line_spacing = 1.04

    r_es_lbl = p4.add_run("ES — ")
    r_es_lbl.bold = True
    r_es_lbl.font.name = FONT_PRIMARY
    r_es_lbl.font.size = Pt(7.5)
    r_es_lbl.font.color.rgb = RGB_ACCENT

    r_es_txt = p4.add_run(step["spanish_translation"])
    r_es_txt.italic = True
    r_es_txt.font.name = FONT_PRIMARY
    r_es_txt.font.size = Pt(7.5)
    r_es_txt.font.color.rgb = RGB_MUTED

    # Line 5: SALES VALUE (WHY BUY FOR FAMILIES)
    p5 = cell.add_paragraph()
    p5.paragraph_format.space_before = Pt(0)
    p5.paragraph_format.space_after = Pt(0)
    p5.paragraph_format.line_spacing = 1.04

    r_val_lbl = p5.add_run("WHY BUY — ")
    r_val_lbl.bold = True
    r_val_lbl.font.name = FONT_PRIMARY
    r_val_lbl.font.size = Pt(7.5)
    r_val_lbl.font.color.rgb = RGB_DARK

    r_val_txt = p5.add_run(step["sales_pitch"])
    r_val_txt.font.name = FONT_PRIMARY
    r_val_txt.font.size = Pt(7.5)
    r_val_txt.font.color.rgb = RGB_BODY


# ==============================================================================
# PHASE BLOCK BUILDER
# ==============================================================================
def add_unified_phase_block(doc, phase_code, phase_title, phase_sub, phase_steps):
    tbl = doc.add_table(rows=4, cols=2)
    force_table_col_widths(tbl, [3.65, 3.65])
    for r in tbl.rows:
        prevent_row_split(r)

    # Row 0: Merged Phase Header
    c_hdr = tbl.cell(0, 0).merge(tbl.cell(0, 1))
    set_explicit_cell_dxa(c_hdr, 7.30)
    set_cell_background(c_hdr, HEX_SURFACE_ALT)
    set_hairline_box_borders(c_hdr, border_color=HEX_BORDER_LIGHT, bottom_accent=HEX_PRIMARY_ACCENT)
    set_cell_margins(c_hdr, top=45, bottom=45, left=95, right=95)

    p = c_hdr.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True

    r0 = p.add_run(f"{phase_code}   /   ")
    r0.bold = True
    r0.font.name = FONT_PRIMARY
    r0.font.size = Pt(8.5)
    r0.font.color.rgb = RGB_ACCENT

    r1 = p.add_run(phase_title)
    r1.bold = True
    r1.font.name = FONT_PRIMARY
    r1.font.size = Pt(9)
    r1.font.color.rgb = RGB_DARK

    r2 = p.add_run(f"   —   {phase_sub}")
    r2.font.name = FONT_PRIMARY
    r2.font.size = Pt(8)
    r2.font.color.rgb = RGB_MUTED

    # Row 1: Step 0 & Step 1
    populate_step_cell(tbl.cell(1, 0), phase_steps[0])
    populate_step_cell(tbl.cell(1, 1), phase_steps[1])
    set_explicit_cell_dxa(tbl.cell(1, 0), 3.65)
    set_explicit_cell_dxa(tbl.cell(1, 1), 3.65)

    # Row 2: Step 2 & Step 3
    populate_step_cell(tbl.cell(2, 0), phase_steps[2])
    populate_step_cell(tbl.cell(2, 1), phase_steps[3])
    set_explicit_cell_dxa(tbl.cell(2, 0), 3.65)
    set_explicit_cell_dxa(tbl.cell(2, 1), 3.65)

    # Row 3: Merged Step 4 (5th step in phase)
    c_last = tbl.cell(3, 0).merge(tbl.cell(3, 1))
    populate_step_cell(c_last, phase_steps[4])
    set_explicit_cell_dxa(c_last, 7.30)


# ==============================================================================
# MAIN BUILDER FUNCTION
# ==============================================================================
def build_present_perfect_infographic_docx(showcase_path, flow_path, out_docx_path):
    doc = Document()
    configure_modern_word_settings(doc)

    style_normal = doc.styles["Normal"]
    style_normal.font.name = FONT_PRIMARY
    style_normal.font.size = Pt(8.5)
    style_normal.font.color.rgb = RGB_BODY
    rPr = style_normal.element.get_or_add_rPr()
    rPr.append(parse_xml(f'<w:lang {nsdecls("w")} w:val="en-US" w:eastAsia="en-US" w:bidi="en-US"/>'))

    for s in doc.sections:
        s.page_width = Inches(8.5)
        s.page_height = Inches(11.0)
        s.top_margin = Inches(0.40)
        s.bottom_margin = Inches(0.40)
        s.left_margin = Inches(0.60)
        s.right_margin = Inches(0.60)

        header = s.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hr = hp.add_run("MEDREMINDER AI   /   CONSUMER SALES PITCH INFOGRAPHIC")
        hr.font.name = FONT_PRIMARY
        hr.font.size = Pt(7.5)
        hr.font.color.rgb = RGB_MUTED

        footer = s.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fr1 = fp.add_run("Sales Pitch Infographic  ·  15 Short Present Perfect Verbs  ·  Página ")
        fr1.font.name = FONT_PRIMARY
        fr1.font.size = Pt(7.5)
        fr1.font.color.rgb = RGB_MUTED
        fp._p.append(parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls("w")))
        fr2 = fp.add_run(" de ")
        fr2.font.name = FONT_PRIMARY
        fr2.font.size = Pt(7.5)
        fr2.font.color.rgb = RGB_MUTED
        fp._p.append(parse_xml(r'<w:fldSimple %s w:instr="NUMPAGES"/>' % nsdecls("w")))

    # ==========================================================================
    # HERO BANNER + 4 KPI METRICS TABLE (ROW 0 & ROW 1)
    # ==========================================================================
    tbl_top = doc.add_table(rows=4, cols=4)
    force_table_col_widths(tbl_top, [1.825, 1.825, 1.825, 1.825])
    for r in tbl_top.rows:
        prevent_row_split(r)

    c_hero = tbl_top.cell(0, 0).merge(tbl_top.cell(0, 3))
    set_explicit_cell_dxa(c_hero, 7.30)
    set_cell_background(c_hero, HEX_PRIMARY_DARK)
    set_hairline_box_borders(c_hero, border_color=HEX_PRIMARY_DARK)
    set_cell_margins(c_hero, top=75, bottom=75, left=110, right=110)

    ph1 = c_hero.paragraphs[0]
    ph1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    ph1.paragraph_format.space_before = Pt(0)
    ph1.paragraph_format.space_after = Pt(2)
    rh_kicker = ph1.add_run("COMMERCIAL PRODUCT LAUNCH & SALES INFOGRAPHIC")
    rh_kicker.bold = True
    rh_kicker.font.name = FONT_PRIMARY
    rh_kicker.font.size = Pt(7.5)
    rh_kicker.font.color.rgb = RGBColor(56, 189, 248)  # Sky 400

    ph2 = c_hero.add_paragraph()
    ph2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    ph2.paragraph_format.space_before = Pt(0)
    ph2.paragraph_format.space_after = Pt(2)
    rh_title = ph2.add_run("MedReminder AI: The Smart Healthcare Assistant for Families")
    rh_title.bold = True
    rh_title.font.name = FONT_PRIMARY
    rh_title.font.size = Pt(14.5)
    rh_title.font.color.rgb = RGBColor(255, 255, 255)

    ph3 = c_hero.add_paragraph()
    ph3.alignment = WD_ALIGN_PARAGRAPH.LEFT
    ph3.paragraph_format.space_before = Pt(0)
    ph3.paragraph_format.space_after = Pt(0)
    rh_sub = ph3.add_run(
        "A 15-point consumer sales pitch explaining why MedReminder AI was created and named, written with simple, high-impact Present Perfect action verbs."
    )
    rh_sub.font.name = FONT_PRIMARY
    rh_sub.font.size = Pt(8.5)
    rh_sub.font.color.rgb = RGB_LIGHT

    kpis = [
        ("15 SHORT VERBS", "Present Perfect (have/has + V3)"),
        ("NAME ORIGIN", "Step 01: Med + Reminder + AI"),
        ("ZERO REPETITION", "15 Unique Action Verbs"),
        ("100% SALES PITCH", "For Families, Seniors & Caregivers"),
    ]
    for idx, (k_big, k_sm) in enumerate(kpis):
        ck = tbl_top.cell(1, idx)
        set_explicit_cell_dxa(ck, 1.825)
        set_cell_background(ck, HEX_SURFACE_SOFT)
        set_hairline_box_borders(ck, border_color=HEX_BORDER_LIGHT, bottom_accent=HEX_PRIMARY_ACCENT)
        set_cell_margins(ck, top=40, bottom=40, left=60, right=60)
        pk = ck.paragraphs[0]
        pk.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pk.paragraph_format.space_before = Pt(0)
        pk.paragraph_format.space_after = Pt(0)
        rk1 = pk.add_run(k_big + "\n")
        rk1.bold = True
        rk1.font.name = FONT_PRIMARY
        rk1.font.size = Pt(8)
        rk1.font.color.rgb = RGB_DARK
        rk2 = pk.add_run(k_sm)
        rk2.font.name = FONT_PRIMARY
        rk2.font.size = Pt(7)
        rk2.font.color.rgb = RGB_MUTED

    # Row 2: Pillars Flow Image
    c_flow = tbl_top.cell(2, 0).merge(tbl_top.cell(2, 3))
    set_explicit_cell_dxa(c_flow, 7.30)
    set_cell_background(c_flow, "FFFFFF")
    set_hairline_box_borders(c_flow, border_color=HEX_BORDER_LIGHT)
    set_cell_margins(c_flow, top=30, bottom=20, left=20, right=20)
    p_flow = c_flow.paragraphs[0]
    p_flow.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_flow.paragraph_format.space_before = Pt(0)
    p_flow.paragraph_format.space_after = Pt(0)
    if os.path.exists(flow_path):
        p_flow.add_run().add_picture(flow_path, width=Inches(7.15))

    # Row 3: Product Showcase Banner
    c_show = tbl_top.cell(3, 0).merge(tbl_top.cell(3, 3))
    set_explicit_cell_dxa(c_show, 7.30)
    set_cell_background(c_show, "FFFFFF")
    set_hairline_box_borders(c_show, border_color=HEX_BORDER_LIGHT)
    set_cell_margins(c_show, top=20, bottom=30, left=20, right=20)
    p_show = c_show.paragraphs[0]
    p_show.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_show.paragraph_format.space_before = Pt(0)
    p_show.paragraph_format.space_after = Pt(0)
    if os.path.exists(showcase_path):
        p_show.add_run().add_picture(showcase_path, width=Inches(7.15))

    add_micro_spacer(doc, pt_size=5)

    # ==========================================================================
    # PILLAR 01, PILLAR 02, PILLAR 03 (UNIFIED TABLES)
    # ==========================================================================
    p1_steps = [s for s in VERBS_PRESENT_PERFECT_DATA if s["phase"] == 1]
    add_unified_phase_block(
        doc,
        "PILLAR 01",
        "NAME ORIGIN & EASY EVERYDAY USE (STEPS 01–05)",
        "Step 01 Explains the Brand Name; Steps 02–05 Make Taking Pills Effortless",
        p1_steps,
    )

    add_micro_spacer(doc, pt_size=6)

    p2_steps = [s for s in VERBS_PRESENT_PERFECT_DATA if s["phase"] == 2]
    add_unified_phase_block(
        doc,
        "PILLAR 02",
        "FAMILY PRIVACY & TIMELY HEALTH ALERTS (STEPS 06–10)",
        "100% Confidential Data, Zero App Freezes & Timely Pharmacy Refills",
        p2_steps,
    )

    add_micro_spacer(doc, pt_size=6)

    p3_steps = [s for s in VERBS_PRESENT_PERFECT_DATA if s["phase"] == 3]
    add_unified_phase_block(
        doc,
        "PILLAR 03",
        "FAMILY TRUST & LIFE-SAVING VALUE (STEPS 11–15)",
        "Instant Alerts, Doctor-Ready Reports & Peace of Mind for Loved Ones",
        p3_steps,
    )

    add_micro_spacer(doc, pt_size=8)

    # ==========================================================================
    # GRAMMAR VERIFICATION & SALES MATRIX TABLE (15 ULTRA-SHORT VERBS)
    # ==========================================================================
    tbl_audit = doc.add_table(rows=18, cols=6)
    col_w = [0.60, 1.00, 1.00, 1.30, 2.00, 1.40]
    force_table_col_widths(tbl_audit, col_w)
    set_minimal_table_borders(tbl_audit)
    for r in tbl_audit.rows:
        prevent_row_split(r)

    # Row 0: Merged Banner
    c_idx = tbl_audit.cell(0, 0).merge(tbl_audit.cell(0, 5))
    set_explicit_cell_dxa(c_idx, 7.30)
    set_cell_background(c_idx, HEX_SURFACE_ALT)
    set_hairline_box_borders(c_idx, border_color=HEX_BORDER_LIGHT, bottom_accent=HEX_PRIMARY_ACCENT)
    set_cell_margins(c_idx, top=45, bottom=45, left=95, right=95)
    pi = c_idx.paragraphs[0]
    pi.paragraph_format.space_before = Pt(0)
    pi.paragraph_format.space_after = Pt(0)
    pi.paragraph_format.keep_with_next = True
    ri0 = pi.add_run("GRAMMAR MATRIX   /   ")
    ri0.bold = True
    ri0.font.name = FONT_PRIMARY
    ri0.font.size = Pt(8.5)
    ri0.font.color.rgb = RGB_ACCENT
    ri1 = pi.add_run("15 SHORT VERBS IN PRESENT PERFECT (HAVE / HAS + V3)")
    ri1.bold = True
    ri1.font.name = FONT_PRIMARY
    ri1.font.size = Pt(9)
    ri1.font.color.rgb = RGB_DARK
    ri2 = pi.add_run("   —   Consumer Sales Pitch & Verification")
    ri2.font.name = FONT_PRIMARY
    ri2.font.size = Pt(8)
    ri2.font.color.rgb = RGB_MUTED

    # Row 1: Column Headers
    headers = ["STEP", "INFINITIVE", "PARTICIPLE (V3)", "PRESENT PERFECT", "SHORT SALES PITCH", "TRADUCCIÓN"]
    for c_i, h_text in enumerate(headers):
        cell = tbl_audit.cell(1, c_i)
        set_explicit_cell_dxa(cell, col_w[c_i])
        set_cell_background(cell, HEX_PRIMARY_DARK)
        set_cell_margins(cell, top=50, bottom=50, left=60, right=60)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(h_text)
        r.bold = True
        r.font.name = FONT_PRIMARY
        r.font.size = Pt(7.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    # Rows 2..16: 15 Distinct Verbs
    for idx, step in enumerate(VERBS_PRESENT_PERFECT_DATA, start=2):
        row = tbl_audit.rows[idx]
        bg = HEX_SURFACE_SOFT if idx % 2 == 0 else "FFFFFF"
        vals = [
            f"Step {step['num']}",
            f"to {step['base_verb']}",
            step["participle"],
            step["perfect_phrase"],
            f"{step['pre_text']}{step['perfect_phrase']}{step['post_text']}",
            step["spanish_translation"]
        ]
        for c_i, val in enumerate(vals):
            cell = row.cells[c_i]
            set_explicit_cell_dxa(cell, col_w[c_i])
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=35, bottom=35, left=60, right=60)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.name = FONT_PRIMARY
            r.font.size = Pt(7.5)
            if c_i in (0, 3):
                r.bold = True
            if c_i == 3:
                r.font.color.rgb = RGB_VERB
            elif c_i == 4:
                r.font.color.rgb = RGB_DARK
            elif c_i == 5:
                r.font.color.rgb = RGB_MUTED
            else:
                r.font.color.rgb = RGB_BODY

    # Row 17: Merged Commercial Closing & Call to Action (CTA)
    c_sum = tbl_audit.cell(17, 0).merge(tbl_audit.cell(17, 5))
    set_explicit_cell_dxa(c_sum, 7.30)
    set_cell_background(c_sum, HEX_SURFACE_ALT)
    set_card_borders(c_sum, left_accent_hex=HEX_PRIMARY_ACCENT, border_color=HEX_BORDER_LIGHT)
    set_cell_margins(c_sum, top=65, bottom=65, left=95, right=95)

    ps = c_sum.paragraphs[0]
    ps.paragraph_format.space_before = Pt(0)
    ps.paragraph_format.space_after = Pt(0)
    ps.paragraph_format.line_spacing = 1.12

    rs1 = ps.add_run("WHY BUY MEDREMINDER AI TODAY? (CALL TO ACTION) — ")
    rs1.bold = True
    rs1.font.name = FONT_PRIMARY
    rs1.font.size = Pt(8.5)
    rs1.font.color.rgb = RGB_DARK

    rs2 = ps.add_run(
        "MedReminder AI is the simplest, most caring way to protect your family's health. We have chosen a name that says it all: Medicine, Reminders, and Smart Care. With big readable buttons, offline alarms that always ring, and private health records, thousands of families have trusted our app to keep their loved ones safe. Download MedReminder AI today and protect what matters most!"
    )
    rs2.font.name = FONT_PRIMARY
    rs2.font.size = Pt(8.5)
    rs2.font.color.rgb = RGB_BODY

    os.makedirs(os.path.dirname(out_docx_path), exist_ok=True)
    doc.save(out_docx_path)
    print(f"Infographic document saved successfully to: {out_docx_path}")


# ==============================================================================
# MAIN EXECUTION
# ==============================================================================
def main():
    screenshots_dir = "/Users/4strob0y/Developer/MedReminderAI/docs/screenshots"
    assets_dir = "/Users/4strob0y/Developer/MedReminderAI/docs/infographic_assets"

    showcase_path, flow_path = generate_visual_assets(screenshots_dir, assets_dir)

    out_dev = "/Users/4strob0y/Developer/MedReminderAI/docs/MedReminderAI_Infografia_15_Verbos_Presente_Perfecto.docx"
    out_dev_en = "/Users/4strob0y/Developer/MedReminderAI/docs/MedReminderAI_Infographic_15_Verbs_Present_Perfect.docx"
    out_dl = "/Users/4strob0y/Downloads/MedReminderAI/docs/MedReminderAI_Infografia_15_Verbos_Presente_Perfecto.docx"

    build_present_perfect_infographic_docx(showcase_path, flow_path, out_dev)
    shutil.copy2(out_dev, out_dev_en)
    if os.path.exists(os.path.dirname(out_dl)):
        shutil.copy2(out_dev, out_dl)

    print("Generation complete!")


if __name__ == "__main__":
    main()
